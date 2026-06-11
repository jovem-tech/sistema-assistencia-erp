from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "fontes"
OUTPUT_DIR = ROOT / "artefatos"


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF4, 0xF6, 0xF9)
BORDER_GRAY = "D9DEE7"


INLINE_BOLD_RE = re.compile(r"(\*\*[^*]+\*\*)")


@dataclass
class Preset:
    name: str
    body_justified: bool
    body_after: int
    body_line: float
    h1_before: int
    h1_after: int
    h2_before: int
    h2_after: int
    h3_before: int
    h3_after: int
    table_header_fill: str


NARRATIVE_PROPOSAL = Preset(
    name="narrative_proposal",
    body_justified=True,
    body_after=8,
    body_line=1.333,
    h1_before=18,
    h1_after=10,
    h2_before=12,
    h2_after=6,
    h3_before=8,
    h3_after=4,
    table_header_fill="F4F6F9",
)

STANDARD_BUSINESS_BRIEF = Preset(
    name="standard_business_brief",
    body_justified=False,
    body_after=6,
    body_line=1.10,
    h1_before=16,
    h1_after=8,
    h2_before=12,
    h2_after=6,
    h3_before=8,
    h3_after=4,
    table_header_fill="F2F4F7",
)


def ensure_output_dir() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(dxa))


def set_table_width_and_indent(table, width_dxa: int, indent_dxa: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_dxa))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if row is table.rows[0]:
                grid_col = OxmlElement("w:gridCol")
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                width = tc_w.get(qn("w:w")) if tc_w is not None else "0"
                grid_col.set(qn("w:w"), width)
                grid.append(grid_col)
    table._tbl.insert(1, grid)


def set_paragraph_border_bottom(paragraph, color: str = BORDER_GRAY) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def add_page_field(paragraph, field_name: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_name
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_sep, text, fld_char_end])


def configure_section(section, header_text: str) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.text = header_text
    header_run = header_p.runs[0]
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = GRAY

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.text = "Pagina "
    run = footer_p.runs[0]
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY
    add_page_field(footer_p, "PAGE")


def configure_styles(doc: Document, preset: Preset) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(preset.body_after)
    normal.paragraph_format.line_spacing = preset.body_line
    normal.paragraph_format.alignment = (
        WD_ALIGN_PARAGRAPH.JUSTIFY if preset.body_justified else WD_ALIGN_PARAGRAPH.LEFT
    )

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, preset.h1_before, preset.h1_after),
        ("Heading 2", 13, BLUE, preset.h2_before, preset.h2_after),
        ("Heading 3", 12, DARK_BLUE, preset.h3_before, preset.h3_after),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    if "Body Small" not in doc.styles:
        small = doc.styles.add_style("Body Small", WD_STYLE_TYPE.PARAGRAPH)
        small.base_style = doc.styles["Normal"]
        small.font.name = "Calibri"
        small.font.size = Pt(9.5)
        small.font.color.rgb = GRAY
        small.paragraph_format.space_after = Pt(4)

    if "Callout" not in doc.styles:
        callout = doc.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        callout.base_style = doc.styles["Normal"]
        callout.font.name = "Calibri"
        callout.font.size = Pt(10.5)
        callout.font.color.rgb = INK
        callout.paragraph_format.space_after = Pt(0)
        callout.paragraph_format.line_spacing = 1.15


def markdown_lines(path: Path) -> List[str]:
    return path.read_text(encoding="utf-8").splitlines()


def clean_line(text: str) -> str:
    return text.strip().replace("  ", " ")


def parse_title_block(lines: Sequence[str]) -> Tuple[str, str, List[str], int]:
    title = ""
    subtitle = ""
    metadata: List[str] = []
    idx = 0
    if idx < len(lines) and lines[idx].startswith("# "):
        title = lines[idx][2:].strip()
        idx += 1
    while idx < len(lines) and not lines[idx].strip():
        idx += 1
    if idx < len(lines) and lines[idx].startswith("## "):
        subtitle = lines[idx][3:].strip()
        idx += 1
    while idx < len(lines):
        raw = lines[idx].strip()
        if raw == "---":
            idx += 1
            break
        if raw:
            metadata.append(raw)
        idx += 1
    return title, subtitle, metadata, idx


def add_inline_markdown(paragraph, text: str, *, base_bold: bool = False) -> None:
    chunks = INLINE_BOLD_RE.split(text)
    for chunk in chunks:
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(chunk)
            run.bold = base_bold
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = INK


def add_metadata_grid(doc: Document, metadata_lines: Sequence[str]) -> None:
    table = doc.add_table(rows=max(1, (len(metadata_lines) + 1) // 2), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [4680, 4680]
    values = list(metadata_lines) + [""] * ((table.rows.__len__() * 2) - len(metadata_lines))
    idx = 0
    for row in table.rows:
        for cidx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[cidx])
            set_cell_shading(cell, "FFFFFF")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(2)
            row_text = values[idx]
            if row_text:
                p = cell.paragraphs[0]
                p.style = doc.styles["Normal"]
                add_inline_markdown(p, row_text)
            idx += 1
    set_table_width_and_indent(table, 9360, 0)


def build_proposal_cover(doc: Document, title: str, subtitle: str, metadata: Sequence[str]) -> None:
    add = doc.add_paragraph()
    add.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add.paragraph_format.space_after = Pt(8)
    run = add.add_run("[NOME DA EMPRESA]")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = GRAY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(title.upper())
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = INK

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = GRAY

    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc.paragraph_format.space_after = Pt(20)
    r_desc = p_desc.add_run(
        "Modelo de servico recorrente para manutencao, suporte e continuidade operacional"
    )
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(10.5)
    r_desc.font.bold = True
    r_desc.font.color.rgb = GRAY

    add_metadata_grid(doc, metadata)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)
    set_paragraph_border_bottom(spacer)
    doc.add_paragraph()


def add_note_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_width(cell, 9360)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.style = "Callout"
    add_inline_markdown(p, text)
    set_table_width_and_indent(table, 9360, 120)
    doc.add_paragraph()


def parse_blocks(lines: Sequence[str]) -> List[Tuple[str, List[str]]]:
    blocks: List[Tuple[str, List[str]]] = []
    idx = 0
    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()
        if not stripped:
            idx += 1
            continue
        if stripped.startswith("## "):
            blocks.append(("h1", [stripped[3:].strip()]))
            idx += 1
            continue
        if stripped.startswith("### "):
            blocks.append(("h2", [stripped[4:].strip()]))
            idx += 1
            continue
        if stripped.startswith("> "):
            quote_lines = []
            while idx < len(lines) and lines[idx].strip().startswith("> "):
                quote_lines.append(lines[idx].strip()[2:].strip())
                idx += 1
            blocks.append(("quote", [" ".join(quote_lines)]))
            continue
        if re.match(r"^\d+\.\s", stripped):
            items = []
            while idx < len(lines) and re.match(r"^\d+\.\s", lines[idx].strip()):
                items.append(re.sub(r"^\d+\.\s", "", lines[idx].strip()))
                idx += 1
            blocks.append(("numbered", items))
            continue
        if stripped.startswith("- "):
            items = []
            while idx < len(lines) and lines[idx].strip().startswith("- "):
                items.append(lines[idx].strip()[2:].strip())
                idx += 1
            blocks.append(("bullet", items))
            continue
        if stripped == "---":
            idx += 1
            continue

        para_lines = [stripped]
        idx += 1
        while idx < len(lines):
            next_line = lines[idx].strip()
            if not next_line:
                idx += 1
                break
            if (
                next_line.startswith("## ")
                or next_line.startswith("### ")
                or next_line.startswith("> ")
                or next_line.startswith("- ")
                or re.match(r"^\d+\.\s", next_line)
                or next_line == "---"
            ):
                break
            para_lines.append(next_line)
            idx += 1
        blocks.append(("paragraph", [" ".join(para_lines)]))
    return blocks


def add_blocks(doc: Document, blocks: Sequence[Tuple[str, List[str]]], preset: Preset) -> None:
    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph(payload[0], style="Heading 1")
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif kind == "h2":
            p = doc.add_paragraph(payload[0], style="Heading 2")
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif kind == "quote":
            add_note_callout(doc, payload[0])
        elif kind == "bullet":
            for item in payload:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.20
                add_inline_markdown(p, item)
        elif kind == "numbered":
            for item in payload:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.20
                add_inline_markdown(p, item)
        elif kind == "paragraph":
            text = payload[0]
            p = doc.add_paragraph()
            p.style = "Normal"
            p.paragraph_format.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
                if preset.body_justified
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline_markdown(p, text)


def build_doc(source_path: Path, output_name: str, preset: Preset, header_text: str, *, proposal_cover: bool) -> Path:
    lines = markdown_lines(source_path)
    title, subtitle, metadata, start_idx = parse_title_block(lines)

    doc = Document()
    configure_styles(doc, preset)
    section = doc.sections[0]
    configure_section(section, header_text)

    if proposal_cover:
        build_proposal_cover(doc, title, subtitle, metadata)
    else:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_after = Pt(4)
        r_title = p_title.add_run(title.upper())
        r_title.font.name = "Calibri"
        r_title.font.size = Pt(22)
        r_title.font.bold = True
        r_title.font.color.rgb = INK

        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(10)
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(13)
        r_sub.font.color.rgb = GRAY

        if metadata:
            for meta in metadata:
                if meta.startswith("> "):
                    add_note_callout(doc, meta[2:].strip())
                else:
                    p = doc.add_paragraph(style="Body Small")
                    add_inline_markdown(p, meta)

    blocks = parse_blocks(lines[start_idx:])
    add_blocks(doc, blocks, preset)

    output_path = OUTPUT_DIR / output_name
    doc.save(output_path)
    return output_path


def main() -> None:
    ensure_output_dir()
    build_doc(
        SOURCE_DIR / "proposta-comercial-modelo.md",
        "proposta-comercial-modelo.docx",
        NARRATIVE_PROPOSAL,
        "Proposta Comercial | Pacote de Contrato de Manutencao",
        proposal_cover=True,
    )
    build_doc(
        SOURCE_DIR / "contrato-manutencao-modelo.md",
        "contrato-manutencao-modelo.docx",
        STANDARD_BUSINESS_BRIEF,
        "Contrato de Manutencao | Modelo Comercial",
        proposal_cover=False,
    )


if __name__ == "__main__":
    main()
